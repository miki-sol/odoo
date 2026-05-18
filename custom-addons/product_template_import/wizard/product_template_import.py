import base64
import io
import logging

from odoo import _, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# Маппинг значения колонки «Тип» -> технический код product.template.type.
# Принимаем русские названия из шаблона импорта и сами технические коды.
TYPE_MAP = {
    'товары': 'consu',
    'услуги': 'service',
    'комбинированный': 'combo',
    'consu': 'consu',
    'service': 'service',
    'combo': 'combo',
}


class ProductTemplateImportWizard(models.TransientModel):
    _name = 'product.template.import.wizard'
    _description = 'Wizard: import products from XLSX'

    file = fields.Binary(string='Excel file', required=True)
    filename = fields.Char(string='Filename')
    has_header = fields.Boolean(
        string='First row is a header',
        default=True,
        help='Uncheck if your file starts with data on row 1.',
    )
    report = fields.Html(string='Import report', readonly=True)

    def action_import(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_('Please upload an Excel (.xlsx) file.'))

        rows = self._read_xlsx(base64.b64decode(self.file))
        if self.has_header and rows:
            rows = rows[1:]

        created = []
        skipped = []
        vals_list = []
        seen_skus = set()

        for idx, row in enumerate(rows, start=2 if self.has_header else 1):
            name_raw = row[0] if len(row) > 0 else None
            sku_raw = row[1] if len(row) > 1 else None
            type_raw = row[2] if len(row) > 2 else None

            if all(v in (None, '') for v in (name_raw, sku_raw, type_raw)):
                continue  # полностью пустая строка — молча пропускаем

            name = str(name_raw).strip() if name_raw not in (None, '') else ''
            if not name:
                skipped.append((idx, _('name is required')))
                continue

            if sku_raw in (None, ''):
                skipped.append((idx, _('SKU is required')))
                continue
            try:
                sku = int(sku_raw)
            except (TypeError, ValueError):
                skipped.append((idx, _('invalid SKU: %r') % (sku_raw,)))
                continue

            if type_raw in (None, ''):
                skipped.append((idx, _('type is required')))
                continue
            ptype = TYPE_MAP.get(str(type_raw).strip().lower())
            if not ptype:
                skipped.append((idx, _('invalid type: %r') % (type_raw,)))
                continue

            if sku in seen_skus:
                skipped.append((idx, _('duplicate SKU %s in file') % (sku,)))
                continue
            if self.env['product.template'].search_count([('sku', '=', sku)]):
                skipped.append((idx, _('SKU %s already exists') % (sku,)))
                continue

            seen_skus.add(sku)
            vals_list.append({'name': name, 'sku': sku, 'type': ptype})
            created.append((idx, name, sku, ptype))

        if vals_list:
            self.env['product.template'].create(vals_list)

        self.report = self._format_report(created, skipped)
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'context': self.env.context,
        }

    @staticmethod
    def _read_xlsx(data):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise UserError(_('openpyxl is required to parse Excel files.')) from e

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            _logger.exception('Failed to parse xlsx')
            raise UserError(_('Could not read Excel file: %s') % e) from e

        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]

    @staticmethod
    def _format_report(created, skipped):
        parts = [
            '<p><b>%s</b>: %d</p>' % (_('Products created'), len(created)),
        ]
        if created:
            parts.append('<ul>')
            for row_idx, name, sku, ptype in created:
                parts.append(
                    '<li>%s %d: %s (SKU %s, %s)</li>'
                    % (_('row'), row_idx, name, sku, ptype)
                )
            parts.append('</ul>')

        parts.append(
            '<p><b>%s</b>: %d</p>' % (_('Products skipped'), len(skipped))
        )
        if skipped:
            parts.append('<ul>')
            for row_idx, reason in skipped:
                parts.append('<li>%s %d: %s</li>' % (_('row'), row_idx, reason))
            parts.append('</ul>')
        return ''.join(parts)
