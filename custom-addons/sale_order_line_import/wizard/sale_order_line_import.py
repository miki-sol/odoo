import base64
import io
import logging

from odoo import _, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class SaleOrderLineImportWizard(models.TransientModel):
    _name = 'sale.order.line.import.wizard'
    _description = 'Wizard: import sale order lines from XLSX'

    order_id = fields.Many2one(
        'sale.order',
        string='Sale Order',
        required=True,
        default=lambda self: self.env.context.get('active_id'),
        ondelete='cascade',
    )
    file = fields.Binary(string='Excel file', required=True)
    filename = fields.Char(string='Filename')
    has_header = fields.Boolean(
        string='First row is a header',
        default=True,
        help='Uncheck if your file starts with data on row 1.',
    )
    report = fields.Html(string='Import report', readonly=True)

    def action_import(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_('Please upload an Excel (.xlsx) file.'))

        rows = self._read_xlsx(base64.b64decode(self.file))
        if self.has_header and rows:
            rows = rows[1:]

        created = []
        skipped = []

        product_tmpl_model = self.env['product.template']
        line_vals_list = []

        for idx, row in enumerate(rows, start=2 if self.has_header else 1):
            sku_raw = row[0] if len(row) > 0 else None
            qty_raw = row[1] if len(row) > 1 else None

            if sku_raw in (None, '') and qty_raw in (None, ''):
                continue  # fully empty row — silently skip

            try:
                sku = int(sku_raw)
            except (TypeError, ValueError):
                skipped.append((idx, _('invalid SKU: %r') % (sku_raw,)))
                continue

            try:
                qty = float(qty_raw)
            except (TypeError, ValueError):
                skipped.append((idx, _('invalid quantity: %r') % (qty_raw,)))
                continue
            if qty <= 0:
                skipped.append((idx, _('quantity must be > 0 (got %s)') % (qty,)))
                continue

            tmpl = product_tmpl_model.search([('sku', '=', sku)], limit=1)
            if not tmpl:
                skipped.append((idx, _('no product with SKU %s') % (sku,)))
                continue

            variants = tmpl.product_variant_ids
            if len(variants) != 1:
                skipped.append((idx, _(
                    'product "%(name)s" (SKU %(sku)s) has %(n)s variants — ambiguous',
                ) % {'name': tmpl.display_name, 'sku': sku, 'n': len(variants)}))
                continue

            line_vals_list.append({
                'order_id': self.order_id.id,
                'product_id': variants.id,
                'product_uom_qty': qty,
            })
            created.append((idx, tmpl.display_name, qty))

        if line_vals_list:
            self.env['sale.order.line'].create(line_vals_list)

        self.report = self._format_report(created, skipped)
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'context': self.env.context,
        }

    @staticmethod
    def _read_xlsx(data):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise UserError(_('openpyxl is required to parse Excel files.')) from e

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            _logger.exception('Failed to parse xlsx')
            raise UserError(_('Could not read Excel file: %s') % e) from e

        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]

    @staticmethod
    def _format_report(created, skipped):
        parts = []
        parts.append(
            '<p><b>%s</b>: %d</p>' % (_('Lines created'), len(created))
        )
        if created:
            parts.append('<ul>')
            for row_idx, name, qty in created:
                parts.append(
                    '<li>%s %d: %s &mdash; %s</li>'
                    % (_('row'), row_idx, name, qty)
                )
            parts.append('</ul>')

        parts.append(
            '<p><b>%s</b>: %d</p>' % (_('Lines skipped'), len(skipped))
        )
        if skipped:
            parts.append('<ul>')
            for row_idx, reason in skipped:
                parts.append('<li>%s %d: %s</li>' % (_('row'), row_idx, reason))
            parts.append('</ul>')
        return ''.join(parts)
